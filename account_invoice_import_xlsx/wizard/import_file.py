# Copyright 2021-26 Zeroincombenze srls <https://www.zeroincombenze.it>
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl).
import base64
import datetime
from io import BytesIO

from markupsafe import Markup, escape
from openpyxl import load_workbook
from unidecode import unidecode

from odoo import Command, _, fields, models
from odoo.exceptions import UserError

# Excel column title -> internal name
TNL = {
    "Denominazione Cliente": "partner_id",
    "Partita Iva": "vat",
    "Codice Fiscale": "fiscalcode",
    "Codice": "default_code",
    "Nome": "name",
    "Quantita": "quantity",
    "Imponibile": "price_subtotal",
    "Iva": "price_tax",
    "Numero": "move_name",
    "Data": "invoice_date",
}
TAX_RATES = (0, 4, 5, 10, 22)


class ImportState:
    """Working data of one import run (records accept no extra attributes)"""

    def __init__(self, company, sale_purchase):
        self.numrow = 0
        self.company = company
        self.sale_purchase = sale_purchase
        self.taxes = {}
        self.sequence = None
        self.invoice_vals = {}
        self.line_vals = {}
        self.partner_domain = []
        self.partner_search_by = []
        self.partner_vals = {}
        self.product_domain = []
        self.product_search_by = []
        self.product_vals = {}


class WizardImportInvoiceFileXlsx(models.TransientModel):
    _name = "wizard.import.invoice.file.xlsx"
    _description = "Import invoice lines from xlsx"

    def _default_journal(self):
        return self.env["account.journal"].search(
            [("type", "=", "sale"), ("company_id", "=", self.env.company.id)],
            limit=1,
        )

    data_file = fields.Binary(string="Excel Data File")
    filename = fields.Char()
    journal_id = fields.Many2one(
        "account.journal",
        string="Journal",
        required=True,
        default=_default_journal,
        domain=[("type", "in", ["sale", "purchase"])],
    )
    invoice_date = fields.Date(string="Invoice Date")
    dry_run = fields.Boolean(string="Dry-run", default=False)
    tracelog = fields.Html("Result History", readonly=True)

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------
    @staticmethod
    def _html_row(*cells):
        return Markup("<tr>%s</tr>\n") % Markup("").join(
            Markup("<td>%s</td>") % escape(cell if cell is not None else "")
            for cell in cells
        )

    @staticmethod
    def get_dim_text(name):
        """Build an ilike pattern tolerant to accents and punctuation"""
        stext = ""
        for ch in name or "":
            if ch.isalpha():
                stext += "_" if unidecode(ch) != ch else ch
            elif not stext.endswith("%"):
                stext += "%"
        return stext

    @staticmethod
    def get_iso_date(value):
        if isinstance(value, datetime.datetime):
            return value.date()
        if isinstance(value, datetime.date):
            return value
        if isinstance(value, str):
            value = value.strip()
            for sep in ("/", "-", "."):
                if sep in value:
                    items = [int(x) for x in value.split(sep)]
                    if len(items) == 3:
                        if items[0] > 31:  # yyyy-mm-dd
                            return datetime.date(items[0], items[1], items[2])
                        return datetime.date(items[2], items[1], items[0])
        return value or False

    def get_float_eval(self, value):
        """Italian formatted number ("2.049,18") or Excel number -> float"""
        if isinstance(value, str):
            value = value.strip().replace(".", "").replace(",", ".")
            try:
                return float(value) if value else 0.0
            except ValueError as e:
                raise UserError(_("Invalid number %s") % value) from e
        return float(value or 0.0)

    def get_data(self):
        if not self.data_file:
            raise UserError(_("No Excel file supplied."))
        wb = load_workbook(BytesIO(base64.b64decode(self.data_file)), data_only=True)
        sheet = wb.active
        rows = sheet.iter_rows(values_only=True)
        colnames = next(rows, ())
        contents = []
        for line in rows:
            if not any(cell not in (None, "") for cell in line):
                continue
            contents.append(
                {
                    colnames[col]: value
                    for col, value in enumerate(line)
                    if colnames[col]
                }
            )
        return contents

    def _fiscalcode_field(self):
        partner_fields = self.env["res.partner"]._fields
        for name in ("l10n_it_codice_fiscale", "fiscalcode"):
            if name in partner_fields:
                return name
        return False

    # ------------------------------------------------------------------
    # Partner / product lookup
    # ------------------------------------------------------------------
    def get_partner(self, st, search_by, disable_err=False):
        html = Markup("")
        partner_model = self.env["res.partner"]
        domain = list(st.partner_domain) + [("type", "=", "contact")]
        stext = self.get_dim_text(st.partner_vals.get("name"))
        cf_field = self._fiscalcode_field()
        if search_by == "vat,fiscalcode":
            domain.append(("vat", "=", st.partner_vals["vat"]))
            if cf_field:
                domain.append((cf_field, "=", st.partner_vals["fiscalcode"]))
        elif search_by == "fiscalcode":
            if not cf_field:
                return False, html
            domain.append((cf_field, "=", st.partner_vals["fiscalcode"]))
        elif search_by != "name":
            domain.append((search_by, "=", st.partner_vals[search_by]))
        if stext:
            domain.append(("name", "ilike", stext))
        partners = partner_model.search(domain)
        if not partners and search_by != "name" and stext:
            partners = partner_model.search(domain[:-1])
        mesg = False
        if len(partners) > 1:
            mesg = _("Found multiple records.")
        elif not partners and not disable_err:
            mesg = _("No partner found.")
        if mesg:
            html = self._html_row(
                st.numrow,
                "",
                st.partner_vals.get("name", ""),
                st.partner_vals.get("vat", ""),
                mesg,
            )
        return partners[:1], html

    def get_product(self, st, search_by, disable_err=False):
        html = Markup("")
        product_model = self.env["product.product"]
        stext = self.get_dim_text(st.line_vals.get("name"))
        domain = list(st.product_domain)
        domain.append((search_by, "=", st.product_vals[search_by]))
        if stext:
            domain.append(("name", "ilike", stext))
        products = product_model.search(domain)
        if not products and stext:
            products = product_model.search(domain[:-1])
        mesg = False
        if len(products) > 1:
            mesg = _("Found multiple records.")
        elif not products and not disable_err:
            mesg = _("No product found.")
        if mesg:
            html = self._html_row(
                st.numrow,
                st.product_vals.get("default_code", ""),
                st.line_vals.get("name", ""),
                "",
                mesg,
            )
        return products[:1], html

    # ------------------------------------------------------------------
    # Row parsing
    # ------------------------------------------------------------------
    def unpack_data(self, st, row):
        st.invoice_vals = {
            "invoice_date": self.invoice_date,
            "journal_id": self.journal_id.id,
            "move_type": (
                "in_invoice" if self.journal_id.type == "purchase" else "out_invoice"
            ),
        }
        st.line_vals = {"name": _("Generic product")}
        st.partner_domain = []
        st.partner_search_by = []
        st.partner_vals = {}
        st.product_domain = []
        st.product_search_by = []
        st.product_vals = {}
        st.sequence = None
        line_subtotal = 0.0

        for field, value in row.items():
            name = TNL.get(field, field)
            if name == "sequence":
                st.sequence = value
            elif value in (None, ""):
                continue
            elif name == "vat":
                value = str(value).strip()
                st.partner_vals[name] = "IT%s" % value if value.isdigit() else value
                st.partner_search_by.insert(0, "vat")
                if "fiscalcode" in st.partner_search_by:
                    st.partner_search_by.insert(0, "vat,fiscalcode")
            elif name == "fiscalcode":
                st.partner_vals[name] = str(value).strip()
                st.partner_search_by.append("fiscalcode")
                if "vat" in st.partner_search_by:
                    st.partner_search_by.insert(0, "vat,fiscalcode")
            elif name == "partner_id":
                if isinstance(value, str):
                    st.partner_search_by.append("name")
                    st.partner_vals["name"] = value
                elif isinstance(value, int):
                    st.invoice_vals["partner_id"] = value
                    st.partner_search_by = []
            elif name == "customer":
                st.partner_domain.append(("customer_rank", ">", 0))
                st.sale_purchase = "sale"
            elif name == "supplier":
                st.partner_domain.append(("supplier_rank", ">", 0))
                st.sale_purchase = "purchase"
            elif name == "move_name":
                st.invoice_vals["name"] = str(value)
            elif name == "invoice_date":
                st.invoice_vals[name] = self.get_iso_date(value)
            elif name == "default_code":
                st.product_search_by.insert(0, name)
                st.product_vals[name] = str(value)
            elif name == "quantity":
                st.line_vals[name] = self.get_float_eval(value)
            elif name == "name":
                st.line_vals[name] = str(value)
            elif (
                name == "price_subtotal"
                and not st.line_vals.get("quantity")
                and not st.line_vals.get("price_unit")
            ):
                line_subtotal = self.get_float_eval(value)
                st.line_vals["price_unit"] = line_subtotal
                st.line_vals["quantity"] = 1
            elif name == "price_tax" and line_subtotal:
                rate = self.get_float_eval(value) * 100 / line_subtotal
                tax_rate = min(TAX_RATES, key=lambda x: abs(x - rate))
                code = "%s%s" % (
                    tax_rate,
                    "a" if st.sale_purchase == "purchase" else "v",
                )
                if st.taxes.get(code):
                    st.line_vals["tax_ids"] = [Command.set([st.taxes[code]])]

    def unpack_partner(self, st):
        html = Markup("")
        if st.sequence in (0, None) and not st.invoice_vals.get("partner_id"):
            for ii, by in enumerate(st.partner_search_by):
                cont = (ii + 1) < len(st.partner_search_by)
                partner, html = self.get_partner(st, by, disable_err=cont)
                if partner:
                    st.invoice_vals["partner_id"] = partner.id
                    break
        return html

    def unpack_product(self, st):
        product = self.env["product.product"]
        html = Markup("")
        if st.sequence or st.sequence is None:
            for ii, by in enumerate(st.product_search_by):
                cont = (ii + 1) < len(st.product_search_by)
                product, html = self.get_product(st, by, disable_err=cont)
                if product:
                    st.line_vals["product_id"] = product.id
                    break
        return product, html

    def header_exists(self, st):
        return "partner_id" in st.invoice_vals

    def line_exists(self, st):
        return bool(
            st.line_vals.get("name")
            and st.line_vals.get("quantity")
            and (st.line_vals.get("price_unit") or st.line_vals.get("product_id"))
        )

    # ------------------------------------------------------------------
    # Record creation
    # ------------------------------------------------------------------
    def create_invoice(self, st, vals):
        if self.dry_run:
            html = self._html_row(
                st.numrow, "", vals.get("name", ""), "", _("Invoice will be created.")
            )
            return True, html
        return self.env["account.move"].create(vals), Markup("")

    def create_invoice_line(self, st, invoice, vals):
        if self.dry_run:
            return self._html_row(
                st.numrow, "", vals.get("name", ""), "", _("Line will be created.")
            )
        # Product data (account, price, taxes) is computed by the ORM;
        # values read from the file take precedence.
        invoice.write({"invoice_line_ids": [Command.create(vals)]})
        return Markup("")

    def store_tax_ids(self, st):
        st.taxes = {}
        for rate in TAX_RATES:
            for tax in self.env["account.tax"].search(
                [
                    ("amount", "=", rate),
                    ("amount_type", "=", "percent"),
                    ("price_include", "=", False),
                    ("company_id", "=", st.company.id),
                ],
                order="sequence,id",
            ):
                suffix = {"sale": "v", "purchase": "a"}.get(tax.type_tax_use)
                if suffix:
                    st.taxes.setdefault("%s%s" % (rate, suffix), tax.id)

    def import_file_xlsx(self):
        self.ensure_one()
        tracelog = Markup("<h3>%s</h3>") % _("Import invoices")
        tracelog += Markup('<table class="table table-sm table-bordered">')
        tracelog += self._html_row(_("Row"), _("Code"), _("Name"), _("Vat"), _("Note"))

        active_ids = (
            self.env.context.get("active_ids")
            if self.env.context.get("active_model") == "account.move"
            else False
        )
        st = ImportState(
            self.journal_id.company_id or self.env.company, self.journal_id.type
        )
        self.store_tax_ids(st)

        datas = self.get_data()
        if active_ids:
            for invoice in self.env["account.move"].browse(active_ids):
                if invoice.state != "draft":
                    raise UserError(
                        _("Lines can be imported only into draft invoices.")
                    )
                if invoice.company_id != st.company:
                    st.company = invoice.company_id
                    self.store_tax_ids(st)
                st.sale_purchase = (
                    "purchase" if invoice.is_purchase_document() else "sale"
                )
                for row in datas:
                    st.numrow += 1
                    self.unpack_data(st, row)
                    product, html = self.unpack_product(st)
                    tracelog += html
                    if self.line_exists(st) or product:
                        st.line_vals.setdefault("quantity", 1.0)
                        tracelog += self.create_invoice_line(st, invoice, st.line_vals)
        else:
            invoice = False
            for row in datas:
                st.numrow += 1
                self.unpack_data(st, row)
                tracelog += self.unpack_partner(st)
                if self.header_exists(st):
                    invoice, html = self.create_invoice(st, st.invoice_vals)
                    tracelog += html
                product, html = self.unpack_product(st)
                tracelog += html
                if invoice and self.line_exists(st):
                    tracelog += self.create_invoice_line(st, invoice, st.line_vals)
        tracelog += Markup("</table>")
        self.tracelog = tracelog
        if not active_ids:
            return {
                "name": _("Import result"),
                "type": "ir.actions.act_window",
                "res_model": self._name,
                "view_mode": "form",
                "res_id": self.id,
                "target": "new",
                "view_id": self.env.ref(
                    "account_invoice_import_xlsx.wizard_import_invoice_result"
                ).id,
            }
        return {"type": "ir.actions.act_window_close"}

    def close_window(self):
        return {"type": "ir.actions.act_window_close"}
